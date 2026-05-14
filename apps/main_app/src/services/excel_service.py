import os
import time
import openpyxl
import logging
from typing import List, Dict, Any


class ExcelService:
    """
    EN: Service to handle reading/writing Excel files.
    VI: Dịch vụ xử lý các thao tác đọc/ghi file Excel.
    """

    @staticmethod
    def load_excel_data(file_path: str, sheet_name: str | int = 0) -> List[Dict[str, Any]]:
        """
        EN: Load data from an Excel file into a list of dictionaries.
        VI: Đọc dữ liệu từ file Excel và chuyển thành danh sách dictionary.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File không tồn tại: {file_path}")

        start_time = time.time()
        logging.info(f"⏳ [TRAP-READ] Bắt đầu mở file: {file_path}")
        try:
            # data_only=True giúp lấy giá trị tính toán cuối cùng thay vì công thức
            wb = openpyxl.load_workbook(file_path, data_only=True)
            logging.info(f"  --> [TRAP-READ] Tải workbook xong: {time.time() - start_time:.2f}s")
            
            if isinstance(sheet_name, int):
                ws = wb.worksheets[sheet_name]
            else:
                if sheet_name not in wb.sheetnames:
                    raise ValueError(f"Sheet '{sheet_name}' không tồn tại.")
                ws = wb[sheet_name]

            data = []
            headers = []
            empty_streak = 0
            
            iter_start = time.time()
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i % 1000 == 0 and i > 0: logging.info(f"      - Đã duyệt {i} dòng... ({time.time() - iter_start:.2f}s)")
                if i == 0:
                    # Lấy Header và loại bỏ các cột trống ảo (do format dư)
                    headers = []
                    empty_streak = 0
                    for j, cell in enumerate(row):
                        if cell is not None and str(cell).strip() != "":
                            headers.append(str(cell).strip())
                            empty_streak = 0
                        else:
                            headers.append(f"Column_{j}")
                            empty_streak += 1
                            if empty_streak > 20: # Dừng nếu gặp 20 cột trống liên tiếp
                                break
                    while headers and (headers[-1].startswith("Column_") or headers[-1] == ""):
                        headers.pop()
                    logging.info(f"📋 Các cột Excel tìm thấy: {headers}")
                else:
                    # Chỉ kiểm tra dữ liệu trong phạm vi số lượng header thực tế
                    actual_row_values = row[:len(headers)]
                    if all(cell is None or str(cell).strip() == "" for cell in actual_row_values):
                        empty_streak += 1
                        # Fail-safe chống treo máy: Dừng nếu gặp 50 dòng trống liên tiếp
                        if empty_streak > 50:
                            break
                        continue
                    
                    empty_streak = 0
                    # Map dữ liệu với Header, ép toàn bộ cell về string hoặc None
                    row_data = {
                        headers[j]: (str(cell).strip() if cell is not None and str(cell).strip() != "" else None)
                        for j, cell in enumerate(actual_row_values)
                    }
                    data.append(row_data)
                    
            wb.close()
            logging.info(f"✅ [TRAP-READ] Hoàn tất đọc {len(data)} dòng. Tổng thời gian: {time.time() - start_time:.2f}s")
            return data
        except Exception as e:
            raise ValueError(f"Lỗi khi đọc file Excel {file_path}: {e}")

    @staticmethod
    def _get_headers(ws) -> List[str]:
        headers = []
        empty_streak = 0
        for col_idx in range(1, ws.max_column + 1):
            cell_val = ws.cell(row=1, column=col_idx).value
            if cell_val is not None and str(cell_val).strip() != "":
                headers.append(str(cell_val).strip())
                empty_streak = 0
            else:
                headers.append("")
                empty_streak += 1
                if empty_streak > 20:
                    break
        while headers and headers[-1] == "":
            headers.pop()
        return headers

    @staticmethod
    def _get_actual_max_row(ws, max_col: int) -> int:
        actual_max_row = 1
        empty_streak = 0
        for row_num in range(2, ws.max_row + 1):
            has_data = False
            for col_idx in range(1, max_col + 1):
                cell_val = ws.cell(row=row_num, column=col_idx).value
                if cell_val is not None and str(cell_val).strip() != "":
                    has_data = True
                    break
            if has_data:
                actual_max_row = row_num
                empty_streak = 0
            else:
                empty_streak += 1
                if empty_streak > 50:
                    break
        return actual_max_row

    @staticmethod
    def save_row(file_path: str, match_cols: List[str], match_value: str, row_data_map: List[tuple], sheet_name: str | int = 0) -> None:
        """
        EN: Update a row if match_value is found, else append a new row.
        VI: Cập nhật dòng nếu tìm thấy match_value, ngược lại thêm dòng mới.
        row_data_map: List of tuples (aliases_list, value)
        """
        start_time = time.time()
        logging.info(f"⏳ [TRAP-SAVE] Bắt đầu ghi file: {file_path}")
        if not os.path.exists(file_path):
            wb = openpyxl.Workbook()
            ws = wb.active
            if isinstance(sheet_name, str):
                ws.title = sheet_name
        else:
            # Lưu ý: Không dùng data_only=True để tránh mất công thức khi lưu
            wb = openpyxl.load_workbook(file_path)
            if isinstance(sheet_name, int):
                ws = wb.worksheets[sheet_name]
            else:
                if sheet_name not in wb.sheetnames:
                    ws = wb.create_sheet(sheet_name)
                ws = wb[sheet_name]

        logging.info(f"  --> [TRAP-SAVE] Tải workbook xong: {time.time() - start_time:.2f}s")
        # Đọc Header từ dòng 1
        headers = ExcelService._get_headers(ws)
            
        is_new_sheet = not any(headers)
        if is_new_sheet:
            for col_idx, (aliases, _) in enumerate(row_data_map, 1):
                ws.cell(row=1, column=col_idx, value=aliases[0])
                headers.append(aliases[0])

        # Tìm cột chứa ID để match (không phân biệt hoa thường)
        match_col_idx = -1
        match_cols_lower = [c.lower() for c in match_cols]
        for idx, header in enumerate(headers):
            if header.lower() in match_cols_lower:
                match_col_idx = idx
                break

        # Tìm dòng cuối cùng có dữ liệu thực sự (tránh lỗi max_row do format làm nhảy dòng)
        actual_max_row = ExcelService._get_actual_max_row(ws, len(headers))
        logging.info(f"  --> [TRAP-SAVE] Tìm actual_max_row ({actual_max_row}) xong: {time.time() - start_time:.2f}s")

        # Tìm xem dòng nào có ID khớp với match_value
        target_row_idx = -1
        if match_col_idx != -1:
            for row_num in range(2, actual_max_row + 1):
                cell_val = ws.cell(row=row_num, column=match_col_idx + 1).value
                if cell_val is not None and str(cell_val).strip().lower() == str(match_value).strip().lower():
                    target_row_idx = row_num
                    break

        # Nếu không tìm thấy, tạo dòng mới ở cuối
        if target_row_idx == -1:
            target_row_idx = actual_max_row + 1

        # Cập nhật dữ liệu vào các ô theo map
        headers_lower = [h.lower() for h in headers]
        
        for aliases, value in row_data_map:
            col_idx = -1
            aliases_lower = [a.lower() for a in aliases]
            for idx, h in enumerate(headers_lower):
                if h in aliases_lower:
                    col_idx = idx + 1
                    break
            
            if col_idx == -1:
                # Không tự động tạo cột mới nếu sheet đã có dữ liệu để tránh sinh rác
                if not is_new_sheet:
                    continue
                else:
                    col_idx = len(headers) + 1
                    ws.cell(row=1, column=col_idx, value=aliases[0])
                    headers.append(aliases[0])
                    headers_lower.append(aliases[0].lower())
                
            ws.cell(row=target_row_idx, column=col_idx, value=value)

        logging.info(f"  --> [TRAP-SAVE] Ghi dữ liệu vào cell xong: {time.time() - start_time:.2f}s")
        ExcelService._reindex_stt(ws, max(actual_max_row, target_row_idx))
        logging.info(f"  --> [TRAP-SAVE] Đánh lại STT xong: {time.time() - start_time:.2f}s")
        wb.save(file_path)
        logging.info(f"  --> [TRAP-SAVE] Lưu file (wb.save) xong: {time.time() - start_time:.2f}s")
        wb.close()
        logging.info(f"✅ [TRAP-SAVE] Hoàn tất ghi file. Tổng thời gian: {time.time() - start_time:.2f}s")

    @staticmethod
    def delete_row(file_path: str, match_cols: List[str], match_value: str, sheet_name: str | int = 0) -> bool:
        """
        EN: Delete a row if match_value is found.
        VI: Xóa dòng nếu tìm thấy match_value.
        """
        if not os.path.exists(file_path):
            return False

        wb = openpyxl.load_workbook(file_path)
        if isinstance(sheet_name, int):
            ws = wb.worksheets[sheet_name]
        else:
            if sheet_name not in wb.sheetnames:
                return False
            ws = wb[sheet_name]

        headers = ExcelService._get_headers(ws)

        match_col_idx = -1
        match_cols_lower = [c.lower() for c in match_cols]
        for idx, header in enumerate(headers):
            if header.lower() in match_cols_lower:
                match_col_idx = idx
                break

        actual_max_row = ExcelService._get_actual_max_row(ws, len(headers))

        if match_col_idx != -1:
            for row_num in range(2, actual_max_row + 1):
                cell_val = ws.cell(row=row_num, column=match_col_idx + 1).value
                if cell_val is not None and str(cell_val).strip().lower() == str(match_value).strip().lower():
                    ws.delete_rows(row_num)
                    ExcelService._reindex_stt(ws, actual_max_row - 1)
                    wb.save(file_path)
                    wb.close()
                    return True
        wb.close()
        return False

    @staticmethod
    def _reindex_stt(ws, actual_max_row: int = -1) -> None:
        """
        EN: Re-calculate static STT (1, 2, 3...) for the sheet.
        VI: Đánh lại Số thứ tự tĩnh (1, 2, 3...) cho sheet.
        """
        header_row = 1
        headers = [h.lower() for h in ExcelService._get_headers(ws)]
        stt_col_idx = -1
        for idx, h in enumerate(headers):
            if h in ["stt", "số tt", "số thứ tự"]:
                stt_col_idx = idx + 1
                break
                
        if stt_col_idx != -1:
            max_r = actual_max_row if actual_max_row > 0 else ExcelService._get_actual_max_row(ws, len(headers))
            current_stt = 1
            for row_num in range(header_row + 1, max_r + 1):
                has_data = False
                for col_idx in range(1, len(headers) + 1):
                    if col_idx != stt_col_idx:
                        val = ws.cell(row=row_num, column=col_idx).value
                        if val is not None and str(val).strip() != "":
                            has_data = True
                            break
                if has_data:
                    ws.cell(row=row_num, column=stt_col_idx, value=current_stt)
                    current_stt += 1
                else:
                    ws.cell(row=row_num, column=stt_col_idx, value=None)