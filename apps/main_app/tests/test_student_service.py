import pytest
import os
import json
import openpyxl
from unittest.mock import patch
from datetime import date
from apps.main_app.src.models.entities import Student
from apps.main_app.src.services.student_service import StudentService


def test_get_students_from_excel_success(tmp_path):
    """
    EN: Test mapping Excel dictionary data to Student objects.
    VI: Kiểm thử chuyển đổi dữ liệu Excel (dict) sang đối tượng Student.
    """
    file_path = os.path.join(tmp_path, "Danh Sach SV.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Mã Sinh viên", "Họ và tên", "Ngày sinh", "Email", "Số điện thoại", "Trạng thái học tập", "Số dư học phí"])
    ws.append(["SV2026001", "Nguyễn Văn A", "15/08/2005", "nva@gmail.com", "0987654321", "Đang học", "5,000,000"])
    wb.save(file_path)

    students = StudentService.get_students_from_excel(file_path)

    # Kiểm tra tính đúng đắn của dữ liệu
    assert len(students) == 1
    
    student = students[0]
    assert isinstance(student, Student)
    assert student.student_id == "SV2026001"
    assert student.full_name == "Nguyễn Văn A"
    assert student.email == "nva@gmail.com"
    assert student.study_status == "Đang học"
    assert student.date_of_birth == date(2005, 8, 15)
    assert student.tuition_balance == 5000000


def test_save_student_to_excel(tmp_path):
    """
    EN: Test saving a Student object to Excel.
    VI: Kiểm thử lưu đối tượng Sinh viên xuống Excel.
    """
    file_path = os.path.join(tmp_path, "Danh_Sach_SV.xlsx")
    student = Student(student_id="SV01", full_name="A", date_of_birth=date(2000, 1, 1), email="a@a.com", phone_number="123", study_status="Đang học", tuition_balance=10)
    StudentService.save_student_to_excel(file_path, student)
    
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "SV01"
    assert ws.cell(row=2, column=2).value == "A"
    wb.close()


def test_scan_form_for_enrichment(tmp_path) -> None:
    """
    EN: Test scanning form for new email/phone numbers.
    VI: Kiểm thử chức năng quét form để phát hiện và làm giàu dữ liệu SĐT/Email.
    """
    form_path = os.path.join(tmp_path, "form.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Form Responses 1"
    ws.append(["Mã SV", "Họ Tên", "Số điện thoại", "Email"])
    ws.append(["SV01", "Nguyễn Văn A", "0901234567", "new@email.com"])
    wb.save(form_path)
    
    # SV01 trong DB hiện tại chỉ có thông tin cũ
    student_db = [
        Student("SV01", "Nguyễn Văn A", date(2000, 1, 1), "old@email.com", "0900000000", "Đang học", 0)
    ]
    
    updates_file = os.path.join(tmp_path, "pending_updates.json")
    with patch("apps.main_app.src.services.student_service.StudentService.get_pending_updates_file", return_value=updates_file):
        count = StudentService.scan_form_for_enrichment(form_path, student_db)
        assert count == 2
        
        assert os.path.exists(updates_file)
        with open(updates_file, "r", encoding="utf-8") as f:
            data = json.load(f)
            assert len(data) == 2