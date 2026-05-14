import os
import openpyxl
import pytest
from unittest.mock import MagicMock
from apps.main_app.src.services.excel_service import ExcelService


def test_load_excel_data_file_not_found():
    """
    EN: Test reading a non-existent Excel file raises FileNotFoundError.
    VI: Kiểm thử đọc file Excel không tồn tại sẽ văng lỗi FileNotFoundError.
    """
    with pytest.raises(FileNotFoundError):
        ExcelService.load_excel_data("non_existent_file_abc.xlsx")


def test_load_excel_data_success(tmp_path):
    """
    EN: Test successfully reading an Excel file.
    VI: Kiểm thử đọc file Excel thành công với dữ liệu mẫu.
    """
    # Tạo file Excel giả lập (mock) trong thư mục tạm
    mock_file_path = os.path.join(tmp_path, "mock_data.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["MSV", "HoTen"])
    ws.append(["SV01", "Nguyen Van A"])
    ws.append(["SV02", "Tran Thi B"])
    wb.save(mock_file_path)

    # Gọi service để test
    data = ExcelService.load_excel_data(mock_file_path)
    
    # Assertions
    assert len(data) == 2
    assert data[0]["MSV"] == "SV01"
    assert data[1]["HoTen"] == "Tran Thi B"


def test_dirty_excel_handling() -> None:
    """
    EN: Test handling of dirty Excel files with massive empty formatting.
    VI: Kiểm thử xử lý file Excel bị format rác (max_row, max_column ảo).
    """
    mock_ws = MagicMock()
    # Giả lập cái bẫy của Excel: báo cáo có 1 triệu dòng và 16 ngàn cột
    mock_ws.max_row = 1048576
    mock_ws.max_column = 16384

    # Tuy nhiên dữ liệu thực tế chỉ nằm ở 3 dòng và 3 cột đầu tiên
    def mock_cell(row, column):
        cell = MagicMock()
        cell.value = "Data" if row <= 3 and column <= 3 else None
        return cell

    mock_ws.cell.side_effect = mock_cell

    # Nếu không có cơ chế cắt tỉa (trim), bài test này sẽ bị treo (Freeze)
    headers = ExcelService._get_headers(mock_ws)
    assert len(headers) == 3
    
    actual_max_row = ExcelService._get_actual_max_row(mock_ws, 3)
    assert actual_max_row == 3