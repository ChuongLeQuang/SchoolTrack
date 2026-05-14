import pytest
import os
import openpyxl
from unittest.mock import patch
from apps.main_app.src.services.class_registration_service import RegistrationService
from apps.main_app.src.models.entities import Student
from datetime import date


def test_parse_accounting_file(tmp_path) -> None:
    """
    EN: Test parsing Google Form/Accounting Excel file.
    VI: Kiểm thử việc đọc và bóc tách dữ liệu từ file Excel Kế toán/Google Form.
    """
    file_path = os.path.join(tmp_path, "form.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Timestamp", "Mã Sinh viên", "Họ & Tên", "Email hay Phone", "Bạn muốn đăng ký học lớp nào"])
    ws.append(["11/05/2026", "25807500161", "Nguyễn Phương Uyên", "ngphuonguyen@gmail.com", "Tiếng Anh Level 2 - Tối 2-4-6 (18:00 - 20:00) - Cơ sở Bình Thạnh [Mã: LNH0526NDB]"])
    wb.save(file_path)

    results = RegistrationService.parse_accounting_file(file_path)
    
    assert len(results) == 1
    assert results[0]["student_id"] == "25807500161"
    assert results[0]["full_name"] == "Nguyễn Phương Uyên"
    assert results[0]["contact_info"] == "ngphuonguyen@gmail.com"
    assert results[0]["class_code"] == "LNH0526NDB"


def test_extract_class_code() -> None:
    """
    EN: Test extracting class code from a string using Regex.
    VI: Kiểm thử việc trích xuất mã lớp từ chuỗi sử dụng Regex.
    """
    text = "Tiếng Anh Level 1 - Tối 2-4-6 (18:00 - 20:00) - Cơ sở Q9 [Mã: LNH0526NEA]"
    assert RegistrationService.extract_class_code(text) == "LNH0526NEA"
    assert RegistrationService.extract_class_code("Không có mã ở đây") == ""


def test_count_registrations_from_file(tmp_path) -> None:
    """
    EN: Test counting registrations from an Excel file.
    VI: Kiểm thử đếm số lượng đăng ký từ file.
    """
    file_path = os.path.join(tmp_path, "form2.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Cột bất kỳ"])
    ws.append(["[Mã: LNH0526NDB]"])
    ws.append(["[Mã: LNH0526NEA]"])
    ws.append(["[Mã: LNH0526NEA]"])
    wb.save(file_path)

    counts, warnings = RegistrationService.count_registrations_from_file(file_path, mode=1, students_db=None)
    
    assert len(counts) == 2
    assert counts["LNH0526NEA"] == 2
    assert counts["LNH0526NDB"] == 1


def test_count_registrations_mode_3_smart(tmp_path) -> None:
    """
    EN: Test Smart mode where it falls back to Email/Phone if MSV is wrong.
    VI: Kiểm thử Chế độ thông minh: Tự dò tìm bằng SĐT nếu MSV bị sai.
    """
    file_path = os.path.join(tmp_path, "form3.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["MSV", "Họ Tên", "Số điện thoại", "Chọn lớp"])
    ws.append(["SAI_MÃ", "Nguyễn Văn A", "0901234567", "[Mã: LNH01]"])
    wb.save(file_path)

    students_db = [Student("SV01", "Nguyễn Văn A", date(2000, 1, 1), "", "0901234567", "Đang học", 0)]
    counts, warnings = RegistrationService.count_registrations_from_file(file_path, mode=3, students_db=students_db)
    
    assert len(counts) == 1
    assert counts["LNH01"] == 1
    assert len(warnings) == 0


@patch("apps.main_app.src.services.student_service.StudentService.save_student_to_excel")
def test_sync_accounting_with_form(mock_save_student, tmp_path) -> None:
    """
    EN: Test reconciling accounting data with registrations.
    VI: Kiểm thử việc đối chiếu dữ liệu kế toán và đăng ký.
    """
    form_path = os.path.join(tmp_path, "form4.xlsx")
    wb_form = openpyxl.Workbook()
    ws_form = wb_form.active
    ws_form.append(["MSV", "Họ Tên", "Email", "Chọn lớp"])
    ws_form.append(["SV01", "A", "a@a.com", "[Mã: CLASS_1]"])
    ws_form.append(["SV02", "B", "b@b.com", "[Mã: CLASS_2]"])
    wb_form.save(form_path)

    acc_path = os.path.join(tmp_path, "acc.xlsx")
    wb_acc = openpyxl.Workbook()
    ws_acc = wb_acc.active
    ws_acc.append(["Mã sinh viên", "Mã lớp", "Số tiền"])
    ws_acc.append(["SV01", "CLASS_1", "5,000,000"])
    ws_acc.append(["SV02", "CLASS_3", "3,000,000"])
    ws_acc.append(["SV03", "CLASS_1", "2,000,000"])
    wb_acc.save(acc_path)

    student_db = [
        Student("SV01", "A", date(2000, 1, 1), "", "", "Đang học", tuition_balance=0),
        Student("SV02", "B", date(2000, 1, 1), "", "", "Đang học", tuition_balance=0),
        Student("SV03", "C", date(2000, 1, 1), "", "", "Đang học", tuition_balance=0),
    ]
    
    results, stats = RegistrationService.sync_accounting_with_form(form_path, acc_path, student_db)
    
    assert len(results) == 3
    assert stats["hop_le"] == 1
    assert stats["lech_khop"] == 2
    assert stats["chua_dong"] == 0
    
    assert student_db[0].tuition_balance == 5000000
    mock_save_student.assert_called_once()