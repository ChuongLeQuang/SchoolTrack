import pytest
import os
import openpyxl
from datetime import date
from apps.main_app.src.models.entities import ClassInfo
from apps.main_app.src.services.class_service import ClassService
from unittest.mock import patch


def test_get_classes_from_excel(tmp_path) -> None:
    file_path = os.path.join(tmp_path, "classes.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Đợt 1"
    ws.append(["Mã Lớp", "Tên Lớp", "Niên Khóa", "Địa điểm", "Phòng học", "Buổi Học", "Lịch Học", "Giờ Học", "Thời gian khai giảng", "Ngày kết thúc khóa", "SL Min", "SL Max", "SL Hiện tại", "Trạng thái"])
    ws.append(["LNH08202501", "Toán Cơ Bản", "2025-2026", "Cơ sở Q9", "Phòng 101", "Sáng", "2-4-6", "08:00 - 10:00", "08/2025", "31/12/2025", 10, 30, 5, "Đang mở"])
    wb.save(file_path)

    classes = ClassService.get_classes_from_excel(file_path, "Đợt 1")
    
    assert len(classes) == 1
    assert classes[0].class_code == "LNH08202501"
    assert classes[0].class_name == "Toán Cơ Bản"
    assert classes[0].location == "Cơ sở Q9"
    assert classes[0].min_students == 10
    assert classes[0].max_students == 30
    assert classes[0].current_students == 5
    assert getattr(classes[0], 'class_status', '') == "Đang mở"
    assert classes[0].end_date == date(2025, 12, 31)


def test_get_classes_from_excel_all_sheets(tmp_path) -> None:
    file_path = os.path.join(tmp_path, "classes_all.xlsx")
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Đợt 1"
    ws1.append(["Mã Lớp"]); ws1.append(["CLASS_A"])
    ws2 = wb.create_sheet("Đợt 2")
    ws2.append(["Mã Lớp"]); ws2.append(["CLASS_B"])
    wb.save(file_path)

    classes = ClassService.get_classes_from_excel(file_path, None)

    assert len(classes) == 2
    class_codes = [c.class_code for c in classes]
    assert "CLASS_A" in class_codes and "CLASS_B" in class_codes


@patch("apps.main_app.src.services.class_wave_metadata_service.WaveMetadataService.save_wave_links")
def test_create_academic_year_file(mock_save_links, tmp_path) -> None:
    file_path = os.path.join(tmp_path, "Quan Ly Lop Hoc 2026.xlsx")
    result = ClassService.create_academic_year_file(file_path)
    
    assert result is True
    assert os.path.exists(file_path)
    
    wb = openpyxl.load_workbook(file_path)
    assert "Đợt 1" in wb.sheetnames
    ws = wb["Đợt 1"]
    assert ws.cell(row=1, column=2).value == "Mã Lớp"
    wb.close()
    
    result_exists = ClassService.create_academic_year_file(file_path)
    assert result_exists is False


@patch("apps.main_app.src.services.class_wave_metadata_service.WaveMetadataService.update_class_in_wave_metadata")
def test_save_class_to_excel(mock_update_json, tmp_path) -> None:
    file_path = os.path.join(tmp_path, "save_class.xlsx")
    cls = ClassInfo(class_code="LNH0825NEA", year_code="2025-2026", class_name="TA", location="Q9", room="", session="Tối", schedule="2-4-6", time_slot="18:00", start_month_year="08/2025", end_date=date(2025, 12, 31), min_students=10, max_students=30, current_students=0, class_status="Dự kiến")
    
    ClassService.save_class_to_excel(file_path, "Đợt 1", cls)
    
    wb = openpyxl.load_workbook(file_path)
    ws = wb["Đợt 1"]
    assert ws.cell(row=2, column=1).value == "LNH0825NEA"
    wb.close()


@patch("apps.main_app.src.services.class_wave_metadata_service.WaveMetadataService.update_class_in_wave_metadata")
def test_delete_class_from_excel(mock_update_json, tmp_path) -> None:
    file_path = os.path.join(tmp_path, "delete_class.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Đợt 1"
    ws.append(["Mã Lớp"])
    ws.append(["LNH0825NEA"])
    wb.save(file_path)

    result = ClassService.delete_class_from_excel(file_path, "2025-2026", "Đợt 1", "LNH0825NEA")
    
    assert result is True
    wb = openpyxl.load_workbook(file_path)
    ws = wb["Đợt 1"]
    assert ws.max_row == 1 # Only header left
    wb.close()


def test_update_registration_counts_deep(tmp_path) -> None:
    file_path = os.path.join(tmp_path, "counts.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Đợt 1"
    ws.append(["Mã Lớp", "SL Hiện tại"])
    ws.append(["LNH0825NEA", 0])
    ws.append(["LNH0825NEB", 2])
    wb.save(file_path)
    
    counts = {"LNH0825NEA": 5, "LNH0825NEB": 3}
    ClassService.update_registration_counts(file_path, "Đợt 1", counts)
    
    wb = openpyxl.load_workbook(file_path)
    ws = wb["Đợt 1"]
    assert ws.cell(row=2, column=2).value == 5
    assert ws.cell(row=3, column=2).value == 3
    wb.close()


def test_get_academic_years(tmp_path) -> None:
    data_dir = os.path.join(tmp_path, "data")
    os.makedirs(data_dir, exist_ok=True)
    with open(os.path.join(data_dir, "Quan Ly Lop Hoc 2025-2026.xlsx"), "w") as f: f.write("")
    with open(os.path.join(data_dir, "Quan Ly Lop Hoc 2024-2025.xlsx"), "w") as f: f.write("")
    with open(os.path.join(data_dir, "Danh Sach SV.xlsx"), "w") as f: f.write("")
    
    with patch("apps.main_app.src.services.class_service.ClassService.get_data_dir", return_value=data_dir):
        years = ClassService.get_academic_years()
        assert len(years) == 2
        assert years[0] == "2025-2026"
        assert years[1] == "2024-2025"